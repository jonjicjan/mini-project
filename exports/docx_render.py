"""Editable print prototype, consuming display-ready snapshot schema version 1.

No ORM reads or attendance/score calculations belong in this renderer. The future
snapshot service owns cutoffs and calculations. Fixed rows preserve leaf alignment;
text that cannot fit is visibly replaced and returned verbatim in fit warnings.
"""

import math
import textwrap
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor, Twips

from .geometry import PAGE_HEIGHT_MM, STUDENT_ROW_MM, mm_to_twips
from .leaves import plan_leaves

LOGO_PATH = Path(__file__).resolve().parent.parent / "static" / "img" / "ksit_logo.jpg"


def _mm(value):
    return Twips(mm_to_twips(str(value)))


class Booklet:
    def __init__(self, snapshot):
        self.snapshot = snapshot
        self.warnings = list(snapshot.get("warnings", []))
        self.document = Document()
        normal = self.document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(9)
        normal.font.color.rgb = RGBColor(0, 0, 0)
        pf = normal.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1
        pf.widow_control = False
        compatibility = self.document.settings.element.find(qn("w:compat"))
        if compatibility is not None:
            for setting in compatibility.findall(qn("w:compatSetting")):
                if setting.get(qn("w:name")) == "compatibilityMode":
                    setting.set(qn("w:val"), "15")
        self.page_count = 0
        self.document.core_properties.title = "Faculty register booklet print sample"
        self.document.core_properties.author = "Faculty Register"

    def page(self, width=210):
        if self.page_count:
            section = self.document.add_section(WD_SECTION_START.NEW_PAGE)
            # Section-break paragraphs must not consume a normal text line.
            paragraph = self.document.paragraphs[-1]
            paragraph.paragraph_format.line_spacing = Pt(1)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run().font.size = Pt(1)
        else:
            section = self.document.sections[0]
        section.page_width = _mm(width)
        section.page_height = _mm(PAGE_HEIGHT_MM)
        section.top_margin = section.bottom_margin = _mm(10)
        section.left_margin = section.right_margin = _mm(5 if width == 105 else 10)
        section.header_distance = section.footer_distance = _mm(4)
        footer = section.footer.paragraphs[0]
        footer.text = (
            "PRINT SAMPLE"
            if self.snapshot["register"]["id"] == "synthetic-register"
            else f"Through {self.snapshot['export_date']} / Revision {self.snapshot['source_revision']}"
        )
        footer.text += " / DRAFT" if self.snapshot["draft"] else ""
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(7)
        self.page_count += 1

    def table(self, widths):
        table = self.document.add_table(rows=0, cols=len(widths))
        table.autofit = False
        for column, width in zip(table.columns, widths, strict=True):
            column.width = _mm(width)
        props = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement(f"w:{edge}")
            for key, value in {"val": "single", "sz": "4", "color": "000000"}.items():
                node.set(qn(f"w:{key}"), value)
            borders.append(node)
        props.append(borders)
        margins = OxmlElement("w:tblCellMar")
        for edge in ("top", "bottom", "left", "right"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:w"), "0" if edge in ("top", "bottom") else "40")
            node.set(qn("w:type"), "dxa")
            margins.append(node)
        props.append(margins)
        return table

    def row(self, table, height):
        row = table.add_row()
        row.height = _mm(height)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, col in zip(row.cells, table.columns, strict=True):
            cell.width = col.width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return row

    def text(self, cell, value, height, *, size=8, bold=False, left=False, label="cell"):
        text = "" if value is None else str(value)
        # Conservative glyph bound, including wide Unicode characters. This is
        # a preflight guard, not a claim to reproduce Word's shaping engine.
        width_pt = cell.width.pt - 4
        # One em bounds Arial's widest Latin glyphs. Simulate word wrapping
        # as well, so whitespace cannot hide an extra line.
        capacity = max(1, int(width_pt / size))
        lines = sum(
            max(
                1,
                len(
                    textwrap.wrap(
                        "".join(
                            "XX" if unicodedata.east_asian_width(c) in "WF" else c for c in line
                        ),
                        width=capacity,
                        break_on_hyphens=False,
                    )
                ),
            )
            for line in text.split("\n")
        )
        if text and lines * size * 1.15 > float(height) * 72 / 25.4 - 2:
            self.warnings.append(f"Page {self.page_count}, {label}: text does not fit: {text}")
            text = "[See fit warning]" if width_pt > 75 else "[!]"
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left else WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.bold = bold
        if text == "A":
            run.font.color.rgb = RGBColor.from_string("FF0000")

    def band(self, table, text, height, *, size=9, bold=False):
        row = self.row(table, height)
        cell = row.cells[0].merge(row.cells[-1])
        self.text(cell, text, height, size=size, bold=bold, label="heading or metadata")

    def institution(self, table, height, half=False):
        if not LOGO_PATH.is_file():
            title = (
                "K S GROUP OF INSTITUTIONS\nK. S. INSTITUTE OF TECHNOLOGY"
                if half
                else "KAMMAVARI SANGHAM (R) - 1952\nK S GROUP OF INSTITUTIONS\n"
                "K. S. INSTITUTE OF TECHNOLOGY"
            )
            self.band(table, title, height, size=8 if half else 10, bold=True)
            return

        row = self.row(table, height)
        cell = row.cells[0].merge(row.cells[-1])
        cell.width = _mm(95 if half else 190)

        inner = cell.add_table(rows=1, cols=2)
        inner.alignment = WD_TABLE_ALIGNMENT.CENTER
        inner.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            '<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
            '<w:insideH w:val="none"/><w:insideV w:val="none"/>'
            '</w:tblBorders>'
        )
        inner._tbl.tblPr.append(borders)

        logo_col_width = 12 if half else 14
        text_col_width = 65 if half else 115
        inner.columns[0].width = _mm(logo_col_width)
        inner.columns[1].width = _mm(text_col_width)

        # Logo cell
        c0 = inner.rows[0].cells[0]
        c0.width = _mm(logo_col_width)
        c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.line_spacing = 1
        img_size = Mm(9 if half else 10.5)
        p0.add_run().add_picture(str(LOGO_PATH), width=img_size, height=img_size)

        # Text cell
        c1 = inner.rows[0].cells[1]
        c1.width = _mm(text_col_width)
        c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = Pt(9.5 if half else 10.5)

        if half:
            r1 = p1.add_run("K S GROUP OF INSTITUTIONS\n")
            r1.font.name = "Arial"
            r1.font.size = Pt(8)
            r1.bold = True
            r2 = p1.add_run("K. S. INSTITUTE OF TECHNOLOGY")
            r2.font.name = "Arial"
            r2.font.size = Pt(9)
            r2.bold = True
        else:
            r1 = p1.add_run("KAMMAVARI SANGHAM (R) - 1952\n")
            r1.font.name = "Arial"
            r1.font.size = Pt(8)
            r1.bold = True
            r2 = p1.add_run("K S GROUP OF INSTITUTIONS\n")
            r2.font.name = "Arial"
            r2.font.size = Pt(8.5)
            r2.bold = True
            r3 = p1.add_run("K. S. INSTITUTE OF TECHNOLOGY")
            r3.font.name = "Arial"
            r3.font.size = Pt(11)
            r3.bold = True

        if len(cell.paragraphs) > 0:
            p_first = cell.paragraphs[0]
            cell._tc.remove(p_first._p)

        if len(cell.paragraphs) > 0:
            p_last = cell.paragraphs[-1]
            p_last.paragraph_format.space_before = Pt(0)
            p_last.paragraph_format.space_after = Pt(0)
            p_last.paragraph_format.line_spacing = Pt(1)
            p_last.add_run().font.size = Pt(1)

    def signatures(self, table, half=False):
        self.band(
            table,
            "\nSignature of HOD"
            if half
            else "\nSignature of Faculty                  Signature of HOD",
            16,
            size=8,
        )

    def lessons(self):
        lessons = self.snapshot["lessons"]
        capacities = [20, 25, 25]
        capacities.extend([25] * math.ceil(max(0, len(lessons) - 70) / 25))
        start = 0
        register = self.snapshot["register"]
        for page_index, capacity in enumerate(capacities):
            self.page()
            table = self.table([12, 94, 25, 25, 34])
            if page_index == 0:
                self.institution(table, 25)
                self.band(table, "LESSON PLAN AND EXECUTION", 10, bold=True)
                self.band(
                    table,
                    f"Branch: {register['branch']}    Semester: {register['semester']}"
                    f"    Section: {register['section']}\nSubject: {register['subject']}"
                    f"    Code: {register['subject_code']}\nFaculty: {register['faculty']}",
                    25,
                )
            else:
                self.band(table, "LESSON PLAN AND EXECUTION", 15, bold=True)
            row = self.row(table, 12)
            for cell, title in zip(
                row.cells,
                ["Class\nNo.", "Topics", "Date\nPlan", "Date\nExecution", "Remarks"],
                strict=True,
            ):
                self.text(cell, title, 12, bold=True)
            for index in range(start, start + capacity):
                lesson = lessons[index] if index < len(lessons) else {}
                values = [
                    index + 1,
                    lesson.get("topic", ""),
                    lesson.get("planned_date", ""),
                    lesson.get("execution_date", ""),
                    lesson.get("remarks", ""),
                ]
                row = self.row(table, 9)
                for col, (cell, value) in enumerate(zip(row.cells, values, strict=True)):
                    self.text(
                        cell,
                        value,
                        9,
                        size=8,
                        left=col in (1, 4),
                        label=f"lesson {index + 1} column {col + 1}",
                    )
            if page_index == len(capacities) - 1:
                self.band(
                    table,
                    "\nSignature of Faculty          Signature of HOD"
                    "          Signature of Principal",
                    16,
                    size=8,
                )
            start += capacity

    def attendance(self, leaf):
        half = leaf.kind == "continuation"
        self.page(105 if half else 210)
        assessment_width = 142.5 / 13
        widths = (
            [8, 30, 57] + [9.5] * 10
            if leaf.kind == "roster"
            else [9.5] * 10
            if half
            else [9.5] * 5 + [assessment_width] * 13
        )
        table = self.table(widths)
        self.institution(table, 25, half)
        r = self.snapshot["register"]
        label = f"Students {leaf.student_start + 1} to {leaf.student_end}"
        self.band(
            table,
            f"{r['academic_year']}    Sem {r['semester']} / {r['section']}\n"
            f"{r['branch']}    {r['subject_code']}\n{r['subject']}\n{label}",
            20,
            size=8,
        )
        sessions = self.snapshot["sessions"]
        selected = [
            sessions[i] if i < len(sessions) else {}
            for i in range(leaf.session_start, leaf.session_end)
        ]
        offset = 3 if leaf.kind == "roster" else 0
        # Body starts at 10 + 25 + 20 + 25 = 80 mm on every panel.
        header_rows = [self.row(table, height) for height in (6, 6, 13)]
        for index, session in enumerate(selected):
            date = session.get("date", "")
            date_label = f"{date[8:10]}\n{date[5:7]}" if date else ""
            for row, value, height in zip(
                header_rows,
                [leaf.session_start + index + 1, session.get("period", ""), date_label],
                [6, 6, 13],
                strict=True,
            ):
                self.text(row.cells[offset + index], value, height, size=7)
        if offset:
            for index, label in enumerate(["Sl.\nNo.", "USN", "Name of the Student"]):
                cell = header_rows[0].cells[index].merge(header_rows[2].cells[index])
                self.text(cell, label, 25, size=8, bold=True)
        if leaf.kind == "assessment":
            labels = self.snapshot["assessments"].get("labels")
            if labels:
                labels = list(labels)
            else:
                labels = [
                    "A1",
                    "A2",
                    "A3",
                    "Assign.\navg.",
                    "Test I\nAtt. %",
                    "Test I\nMarks",
                    "Test II\nAtt. %",
                    "Test II\nMarks",
                    "Test III\nAtt. %",
                    "Test III\nMarks",
                    "Final\nAtt. %",
                    "Final\nMarks",
                    "Overall",
                ]
                if self.snapshot["assessments"].get("weighted_contributions"):
                    labels[3], labels[11] = "Assign.\ncontrib.", "Test\ncontrib."
            maxima = self.snapshot["assessments"]["maxima"]
            maxkeys = [
                "A1",
                "A2",
                "A3",
                "assignment",
                None,
                "T1",
                None,
                "T2",
                None,
                "T3",
                None,
                "test",
                "overall",
            ]
            for i, (label, key) in enumerate(zip(labels, maxkeys, strict=True), 5):
                cell = header_rows[0].cells[i].merge(header_rows[2].cells[i])
                maximum = f"\n/{maxima[key]}" if key else ""
                self.text(cell, label + maximum, 25, size=6.5, bold=True)
        students = self.snapshot["students"]
        for index in range(leaf.student_start, leaf.student_end):
            row = self.row(table, STUDENT_ROW_MM)
            if index >= len(students):
                continue
            student = students[index]
            student_id = student["id"]
            if offset:
                for col, key in enumerate(["serial", "usn", "name"]):
                    self.text(
                        row.cells[col],
                        student[key],
                        STUDENT_ROW_MM,
                        size=7.5,
                        left=col == 2,
                        label=f"student {index + 1} {key}",
                    )
            cells = self.snapshot["attendance"].get(student_id, {}).get("cells", {})
            for col, session in enumerate(selected, offset):
                self.text(
                    row.cells[col],
                    cells.get(session.get("id"), ""),
                    STUDENT_ROW_MM,
                    size=8,
                    label=f"student {index + 1} attendance",
                )
            if leaf.kind == "assessment":
                # Summaries belong only to the final batch, not intermediate sets of 75.
                final_batch = leaf.session_end >= len(sessions)
                marks = self.snapshot["assessments"]["students"].get(student_id, {})
                display = marks.get("display", [""] * 13) if final_batch else [""] * 13
                if len(display) != 13:
                    raise ValueError("Assessment display requires thirteen panel values.")
                for cell, value in zip(row.cells[5:], display, strict=True):
                    self.text(
                        cell, value, STUDENT_ROW_MM, size=7, label=f"student {index + 1} assessment"
                    )
        for key, label in [
            ("absent", "No. of Students Absent"),
            ("present", "No. of Students Present"),
        ]:
            row = self.row(table, 6)
            if offset:
                cell = row.cells[0].merge(row.cells[2])
                self.text(cell, label, 6, size=8, left=True)
            for col, session in enumerate(selected, offset):
                # Totals are per block, supplied by the snapshot service.
                totals = self.snapshot["column_totals"].get(str(leaf.student_start), {})
                self.text(row.cells[col], totals.get(session.get("id"), {}).get(key, ""), 6)
        self.band(table, "Faculty initials", 12, size=8)
        self.signatures(table, half)

    def timetable(self):
        timetable = self.snapshot["timetable"]
        entries = timetable.get("entries", [])
        slots = timetable["slots"]
        days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
        valid_periods = {slot["id"] for slot in slots if not slot.get("break")}
        if len({slot["id"] for slot in slots}) != len(slots):
            raise ValueError("Timetable slot IDs must be unique.")
        for entry in entries:
            if entry["day"] not in days or entry["period"] not in valid_periods:
                raise ValueError("A timetable entry has no matching day or teaching slot.")
        if not any(entry["day"] == "Sunday" for entry in entries):
            days.remove("Sunday")
        for start in range(0, max(1, len(slots)), 9):
            selected = slots[start : start + 9]
            if not selected:
                raise ValueError("At least one timetable slot is required.")
            self.page()
            table = self.table([22] + [168 / len(selected)] * len(selected))
            self.institution(table, 25)
            self.band(table, "FACULTY TIMETABLE" + (" continued" if start else ""), 12, bold=True)
            self.band(
                table,
                f"Faculty: {self.snapshot['register']['faculty']}\n"
                f"Effective from: {timetable.get('effective_from', '')}",
                18,
            )
            row = self.row(table, 22)
            self.text(row.cells[0], "Day /\nTime", 22, bold=True)
            for cell, slot in zip(row.cells[1:], selected, strict=True):
                self.text(
                    cell,
                    f"{slot['label']}\n{slot['start']}\n{slot['end']}",
                    22,
                    size=7,
                    bold=True,
                    label="timetable slot",
                )
            day_rows = []
            for day in days:
                row = self.row(table, 18)
                day_rows.append(row)
                self.text(row.cells[0], day[:3].upper(), 18, bold=True)
                for cell, slot in zip(row.cells[1:], selected, strict=True):
                    matches = [
                        entry
                        for entry in entries
                        if entry["day"] == day and entry["period"] == slot["id"]
                    ]
                    value = "\n".join(
                        f"{entry['subject_code']}\n{entry['section']}\n{entry['activity']}"
                        for entry in matches
                    )
                    self.text(cell, value, 18, size=7, label=f"timetable {day} {slot['label']}")
            for index, slot in enumerate(selected, 1):
                if slot.get("break"):
                    cell = day_rows[0].cells[index].merge(day_rows[-1].cells[index])
                    self.text(cell, slot["label"], 18 * len(days), size=8)
            r = self.snapshot["register"]
            self.band(
                table,
                f"Subject: {r['subject']}    Code: {r['subject_code']}\n"
                f"Branch: {r['branch']}    Semester: {r['semester']} / {r['section']}",
                18,
            )
            self.band(
                table,
                "\nSignature of Faculty          Signature of HOD          Signature of Principal",
                16,
                size=8,
            )


def render_booklet(snapshot: dict, target: Path) -> list[str]:
    if snapshot.get("schema_version") != 1:
        raise ValueError("Unsupported snapshot schema version; expected 1.")
    required = {
        "register",
        "export_date",
        "source_revision",
        "students",
        "lessons",
        "sessions",
        "attendance",
        "column_totals",
        "assessments",
        "timetable",
        "warnings",
        "draft",
    }
    if missing := required - snapshot.keys():
        raise ValueError(f"Missing snapshot fields: {', '.join(sorted(missing))}")
    booklet = Booklet(snapshot)
    booklet.lessons()
    for leaf in plan_leaves(len(snapshot["students"]), len(snapshot["sessions"])):
        booklet.attendance(leaf)
    booklet.timetable()
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    booklet.document.save(target)
    return list(dict.fromkeys(booklet.warnings))
